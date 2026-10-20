import pytest
from docx import Document

from products.humanizer.service.document_parser import (
    export_docx_revisions,
    parse_document,
)


def test_docx_keeps_styles_tables_and_complex_paragraphs(tmp_path):
    src = tmp_path / "original.docx"
    out = tmp_path / "revised.docx"
    document = Document()
    document.add_heading("Chapter One", 1)
    simple = document.add_paragraph()
    run = simple.add_run("The original wording is unclear.")
    run.bold = True

    complex_p = document.add_paragraph()
    complex_p.add_run("Mixed ")
    complex_p.add_run("formatting.").italic = True

    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Table content unchanged."
    document.save(src)

    parsed = parse_document(src)
    simple_block = next(x for x in parsed.blocks if x.text == "The original wording is unclear.")
    complex_block = next(x for x in parsed.blocks if x.text == "Mixed formatting.")
    table_block = next(x for x in parsed.blocks if x.text == "Table content unchanged.")

    assert simple_block.editable
    assert complex_block.editable
    assert table_block.kind == "table_cell"

    export_docx_revisions(src, out, {simple_block.index: "The wording is now clearer."})
    original = Document(src)
    revised = Document(out)

    assert original.paragraphs[1].text == "The original wording is unclear."
    assert revised.paragraphs[0].style.name == "Heading 1"
    assert revised.paragraphs[1].runs[0].bold
    assert revised.paragraphs[1].text == "The wording is now clearer."
    assert revised.paragraphs[2].text == "Mixed formatting."
    assert revised.tables[0].cell(0, 0).text == "Table content unchanged."

    export_docx_revisions(src, tmp_path / "mixed.docx",
                          {complex_block.index: "Mixed revised formatting."})
    revised_complex = Document(tmp_path / "mixed.docx").paragraphs[2]
    assert revised_complex.text == "Mixed revised formatting."
    assert revised_complex.runs[1].italic is True
    with pytest.raises(ValueError, match="overwrite"):
        export_docx_revisions(src, src, {})
    with pytest.raises(FileExistsError):
        export_docx_revisions(src, out, {})



def test_text_only_heading_and_table_cell_can_be_edited(tmp_path):
    source, output = tmp_path / "headings.docx", tmp_path / "edited.docx"
    document = Document()
    document.add_heading("Question One", level=2)
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Answer incomplete."
    document.save(source)
    parsed = parse_document(source)
    heading = next(block for block in parsed.blocks if block.kind == "heading")
    table = next(block for block in parsed.blocks if block.kind == "table_cell")
    assert heading.editable and table.editable
    export_docx_revisions(source, output, {
        heading.index: "Question Two",
        table.index: "Answer clarified.",
    })
    doc = Document(output)
    assert doc.paragraphs[0].text == "Question Two"
    assert doc.paragraphs[0].style.name == "Heading 2"
    assert doc.tables[0].cell(0, 0).text == "Answer clarified."


def test_hyperlinks_and_inline_objects_remain_read_only(tmp_path):
    source = tmp_path / "complex.docx"
    document = Document()
    p = document.add_paragraph("Open ")
    p.add_hyperlink("example", "https://example.com")
    document.save(source)
    block = parse_document(source).blocks[0]
    assert not block.editable


def test_multi_run_unchanged_tokens_keep_style(tmp_path):
    source = tmp_path / "original.docx"
    dest = tmp_path / "revised.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Use this ").bold = True
    para.add_run("incorrect word").italic = True
    para.add_run(" in context.")
    doc.save(source)
    block = parse_document(source).blocks[0]
    assert block.editable
    export_docx_revisions(source, dest, {0: "Use this better word in context."})
    edited = Document(dest).paragraphs[0]
    assert edited.text == "Use this better word in context."
    assert edited.runs[0].bold is True
    assert edited.runs[1].italic is True
    assert Document(source).paragraphs[0].text == "Use this incorrect word in context."
