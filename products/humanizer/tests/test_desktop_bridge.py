from docx import Document
import pytest

from products.humanizer.service import desktop_bridge as bridge


def _sample(tmp_path):
    src = tmp_path / "essay.docx"
    doc = Document()
    doc.add_heading("Background", 1)
    doc.add_paragraph("Smith (2024) reported 67.4% adoption.")
    doc.save(src)
    opened = bridge.process({"action": "open", "path": str(src)})
    return src, opened


def test_open_and_export_docx_reviewed_revision(tmp_path):
    src, opened = _sample(tmp_path)
    editable = next(b for b in opened["blocks"] if b["editable"] and b["kind"] == "paragraph")
    assert opened["fingerprint"] and opened["format"] == "docx"

    dest = tmp_path / "edited.docx"
    result = bridge.process({
        "action": "export", "path": str(src),
        "fingerprint": opened["fingerprint"],
        "destination": str(dest),
        "revisions": {str(editable["index"]): "Smith (2024) documented 67.4% adoption."},
    })
    assert result["revised_count"] == 1
    assert Document(src).paragraphs[1].text == "Smith (2024) reported 67.4% adoption."
    assert Document(dest).paragraphs[1].text == "Smith (2024) documented 67.4% adoption."

    with pytest.raises(ValueError, match="new, nonexistent"):
        bridge.process({
            "action": "export", "path": str(src),
            "fingerprint": opened["fingerprint"],
            "destination": str(dest), "revisions": {},
        })


def test_reject_stale_document(tmp_path):
    src, opened = _sample(tmp_path)
    doc = Document(src)
    doc.paragraphs[1].text = "Someone edited this outside the app."
    doc.save(src)
    with pytest.raises(ValueError, match="changed since opening"):
        bridge.process({
            "action": "export", "path": str(src),
            "fingerprint": opened["fingerprint"],
            "destination": str(tmp_path / "new.docx"),
            "revisions": {},
        })


def test_local_rewrite_bridge_with_stubbed_model(tmp_path, monkeypatch):
    src, opened = _sample(tmp_path)
    editable = next(b for b in opened["blocks"] if b["editable"] and b["kind"] == "paragraph")

    def fake_generate(prompt):
        block = prompt.split("Text:\n", 1)[1].split("\n\nReturn only", 1)[0]
        return block.replace("reported", "documented")

    monkeypatch.setattr(bridge, "_get_model", lambda backend, path: fake_generate)
    result = bridge.process({
        "action": "rewrite", "path": str(src),
        "fingerprint": opened["fingerprint"], "index": editable["index"],
        "mode": "grammar", "backend": "gguf", "model_path": "local.gguf",
    })
    assert result["passed"]
    assert result["proposal"] == "Smith (2024) documented 67.4% adoption."

    with pytest.raises(ValueError, match="Unknown rewrite mode"):
        bridge.process({
            "action": "rewrite", "path": str(src),
            "fingerprint": opened["fingerprint"], "index": editable["index"],
            "mode": "fiction", "backend": "gguf", "model_path": "local.gguf",
        })


def test_plain_text_export_keeps_original(tmp_path):
    source = tmp_path / "notes.md"
    source.write_bytes(b"# Overview\r\n\r\nThis is somewhat awkward.\r\n\r\nClosing note.")
    opened = bridge.process({"action": "open", "path": str(source)})
    target = tmp_path / "notes-revised.md"
    bridge.process({
        "action": "export", "path": str(source),
        "fingerprint": opened["fingerprint"], "destination": str(target),
        "revisions": {"1": "This is clearer."},
    })
    assert source.read_bytes() == b"# Overview\r\n\r\nThis is somewhat awkward.\r\n\r\nClosing note."
    assert target.read_bytes() == b"# Overview\r\n\r\nThis is clearer.\r\n\r\nClosing note."


def test_local_diagnostics_does_not_load_weights_or_documents(tmp_path, monkeypatch):
    bogus = tmp_path / "tiny.gguf"
    bogus.write_bytes(b"GGUF")
    monkeypatch.setattr(
        bridge, "_get_model",
        lambda *args, **kwargs: pytest.fail("Diagnostics must not load the model"),
    )
    checks = bridge.process({"action": "diagnostics", "model_path": str(bogus)})
    assert checks["model_file_found"]
    assert checks["gguf_header_valid"]
    assert not checks["components_present"]
    assert checks["model_bytes"] == 4
    assert "model load and rewrite quality are not tested" in checks["note"]


def test_local_diagnostics_reports_missing_gguf(tmp_path):
    result = bridge.process({
        "action": "diagnostics",
        "model_path": str(tmp_path / "missing.gguf"),
    })
    assert not result["model_file_found"]
    assert not result["components_present"]
    assert result["issues"]



def test_local_diagnostics_detects_missing_frozen_english_dictionary(tmp_path, monkeypatch):
    import wordfreq
    broken = tmp_path / "dummy.gguf"
    broken.write_bytes(b"GGUF")
    def missing(*args, **kwargs):
        raise LookupError("No wordlist 'best' available for language 'en'")
    monkeypatch.setattr(wordfreq, "zipf_frequency", missing)
    result = bridge.process({"action": "diagnostics", "model_path": str(broken)})
    assert result["english_dictionary_installed"] is False
    assert any("spelling dictionary" in message for message in result["issues"])
