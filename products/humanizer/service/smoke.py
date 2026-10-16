"""End-to-end local Qwen smoke test for Silabs Humanizer.

Creates a disposable DOCX, invokes the REAL local model, verifies protected
content, exports to a separate file, and prints the proposed rewrite.
Nothing is downloaded or sent to a cloud service.
"""
from __future__ import annotations

import argparse
from pathlib import Path
import tempfile
import time

from .document_parser import export_docx_revisions, parse_document
from .rewrite import rewrite_text
from .style_profile import StyleProfile


SAMPLE = (
    "Smith (2024) reported that 67.4% of participants adopted the platform. "
    "It is important to note that this underscores the crucial role "
    "technology plays in the modern education landscape."
)


def main() -> int:
    parser = argparse.ArgumentParser(description="Real local Qwen3 Humanizer smoke test")
    parser.add_argument("--backend", choices=("gguf", "transformers"), required=True)
    parser.add_argument("--model", required=True, help="Existing local model path")
    parser.add_argument("--device", default="cpu")
    parser.add_argument("--dtype", default="float32")
    parser.add_argument("--gpu-layers", type=int, default=0)
    args = parser.parse_args()

    from docx import Document

    if args.backend == "gguf":
        from .local_model import GGUFLocalModel
        model = GGUFLocalModel(args.model, gpu_layers=args.gpu_layers, max_new_tokens=256)
    else:
        from .local_model import TransformersLocalModel
        model = TransformersLocalModel(
            args.model, device=args.device, dtype=args.dtype, max_new_tokens=256
        )

    with tempfile.TemporaryDirectory(prefix="silabs-humanizer-") as folder:
        source = Path(folder) / "source.docx"
        output = Path(folder) / "result.docx"
        doc = Document()
        doc.add_heading("Humanizer offline pilot", level=1)
        doc.add_paragraph(SAMPLE)
        doc.save(source)

        parsed = parse_document(source)
        target = next(b for b in parsed.blocks if b.text == SAMPLE)
        started = time.perf_counter()
        revised, review = rewrite_text(SAMPLE, model, StyleProfile(mode="academic"))
        seconds = time.perf_counter() - started

        print("SOURCE:", SAMPLE)
        print("\nPROPOSED:", revised)
        print("\nSECONDS:", round(seconds, 2))
        print("PRESERVATION CHECK:", "PASS" if review.passed else "FAIL")
        if review.issues:
            print("ISSUES:", "; ".join(review.issues))
            return 1

        export_docx_revisions(source, output, {target.index: revised})
        result = Document(output)
        actual = result.paragraphs[1].text
        assert actual == revised, "Round-trip export changed the text"
        assert "Smith (2024)" in actual and "67.4%" in actual
        print("DOCX ROUND TRIP: PASS")
        print("Manual meaning and style review: REQUIRED (not automated)")
        return 0


if __name__ == "__main__":
    raise SystemExit(main())
