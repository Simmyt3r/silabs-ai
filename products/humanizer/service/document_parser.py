"""Structure-aware document import and conservative DOCX export.

DOCX rewriting currently permits only plain, single-run paragraphs; complex
multi-run, hyperlink, image, field and nested-table content remains unchanged.
This deliberately avoids the formatting loss caused by Paragraph.text = ....
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Iterator


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx"}


@dataclass(frozen=True)
class DocumentBlock:
    index: int
    kind: str
    text: str
    style: str | None = None
    editable: bool = True
    reason: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    path: Path
    blocks: list[DocumentBlock]


def parse_plain_text(path: str | Path) -> ParsedDocument:
    source = Path(path)
    if source.suffix.lower() not in {".txt", ".md"}:
        raise ValueError(f"Unsupported plain-text format: {source.suffix}")
    text = source.read_text(encoding="utf-8")
    paragraphs = [part for part in text.split("\n\n") if part.strip()]
    blocks = [
        DocumentBlock(index=i, kind="paragraph", text=paragraph)
        for i, paragraph in enumerate(paragraphs)
    ]
    return ParsedDocument(path=source, blocks=blocks)


def _docx_paragraphs(document) -> Iterator[tuple[object, bool]]:
    # Top-level body order matters: document.paragraphs alone loses tables.
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    seen = set()
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document), False
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        # Merged cells may be reported more than once.
                        if paragraph._p not in seen:
                            seen.add(paragraph._p)
                            yield paragraph, True


def _is_safely_editable(paragraph) -> bool:
    if len(paragraph.runs) != 1 or not paragraph.text.strip():
        return False
    # Preserve paragraph layout, styling and all structured/inline content.
    # Only a single text run with no field, drawing, break, hyperlink, etc.
    from docx.oxml.ns import qn

    if any(element.tag not in {qn("w:pPr"), qn("w:r")} for element in paragraph._p):
        return False
    run = paragraph.runs[0]
    if any(element.tag not in {qn("w:rPr"), qn("w:t")} for element in run._r):
        return False
    return True


def parse_docx(path: str | Path) -> ParsedDocument:
    from docx import Document

    source = Path(path)
    document = Document(str(source))
    blocks = []
    for index, (paragraph, in_table) in enumerate(_docx_paragraphs(document)):
        style = paragraph.style.name if paragraph.style else None
        kind = "table_cell" if in_table else "heading" if (style or "").startswith("Heading") else "paragraph"
        editable = _is_safely_editable(paragraph)
        blocks.append(
            DocumentBlock(
                index=index, kind=kind, text=paragraph.text, style=style,
                editable=editable,
                reason=None if editable else "Complex or empty paragraph: export disabled for this block",
            )
        )
    return ParsedDocument(path=source, blocks=blocks)


def parse_document(path: str | Path) -> ParsedDocument:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".txt", ".md"}:
        return parse_plain_text(source)
    if suffix == ".docx":
        return parse_docx(source)
    raise ValueError(f"Unsupported document format: {suffix}")


def export_docx_revisions(
    source: str | Path,
    destination: str | Path,
    revisions: Mapping[int, str],
) -> Path:
    """Write reviewed replacements to a NEW DOCX; never overwrite source.

    Only text-only single-run paragraphs are allowed. Reject all invalid
    changes before saving so the original document remains untouched.
    """
    from docx import Document

    src = Path(source).resolve()
    dst = Path(destination).resolve()
    if src.suffix.lower() != ".docx" or dst.suffix.lower() != ".docx":
        raise ValueError("DOCX export requires .docx source and destination.")
    if src == dst:
        raise ValueError("Cannot overwrite the source document.")
    if dst.exists():
        raise FileExistsError(f"Destination already exists: {dst}")

    document = Document(str(src))
    paragraphs = list(_docx_paragraphs(document))
    for index, revised_text in revisions.items():
        if not isinstance(index, int) or index < 0 or index >= len(paragraphs):
            raise IndexError(f"Unknown document block index: {index}")
        paragraph = paragraphs[index][0]
        if not _is_safely_editable(paragraph):
            raise ValueError(f"Block {index} has complex formatting; refusing unsafe rewrite.")
        if not isinstance(revised_text, str) or not revised_text.strip():
            raise ValueError(f"Block {index} replacement must be non-empty text.")

    for index, revised_text in revisions.items():
        paragraph = paragraphs[index][0]
        paragraph.runs[0].text = revised_text

    dst.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(dst))
    return dst
